"""Tests for the shared HCAI-style multi-sheet workbook parser.

Builds small synthetic workbooks in-test (openpyxl) that reproduce the
real structural quirks HCAI files have: a title row above the header,
footnote rows, and suppression markers -- rather than relying only on
mocks, since the header-row-detection logic is exactly what would break
silently on a malformed real file.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from scc_health_pipeline.sources.workbook import (
    WorkbookParseError,
    find_header_row,
    list_sheet_names,
    parse_all_sheets,
    parse_sheet,
)


def _build_workbook_with_title_row(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ED Encounters"
    ws.append(["HCAI Emergency Department Encounters -- Santa Clara County"])
    ws.append([])
    ws.append(["Facility", "Year", "Encounters", "Payer"])
    ws.append(["Facility A", 2024, 1200, "Medi-Cal"])
    ws.append(["Facility B", 2024, "*", "Medicare"])  # suppressed cell
    ws.append(["Source: HCAI. Suppressed values (*) protect confidentiality."])
    wb.save(path)


def _build_multi_sheet_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["ID", "Value"])
    ws1.append([1, 100])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Code", "Count"])
    ws2.append(["A", 5])
    wb.save(path)


def test_find_header_row_skips_title_and_blank_rows(tmp_path: Path) -> None:
    path = tmp_path / "ed_encounters.xlsx"
    _build_workbook_with_title_row(path)
    header_row = find_header_row(path, "ED Encounters")
    assert header_row == 2  # 0-indexed: title=0, blank=1, header=2


def test_parse_sheet_detects_suppressed_cells(tmp_path: Path) -> None:
    path = tmp_path / "ed_encounters.xlsx"
    _build_workbook_with_title_row(path)
    parsed = parse_sheet(path, "ED Encounters")
    assert list(parsed.dataframe.columns) == ["Facility", "Year", "Encounters", "Payer"]
    assert parsed.suppressed_cell_count == 1
    assert len(parsed.dataframe) == 2  # data rows only, not the footnote row


def test_list_sheet_names(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    _build_multi_sheet_workbook(path)
    assert list_sheet_names(path) == ["Sheet1", "Sheet2"]


def test_parse_all_sheets_parses_every_sheet(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    _build_multi_sheet_workbook(path)
    results = parse_all_sheets(path)
    assert set(results.keys()) == {"Sheet1", "Sheet2"}
    assert list(results["Sheet1"].dataframe.columns) == ["ID", "Value"]
    assert list(results["Sheet2"].dataframe.columns) == ["Code", "Count"]


def test_find_header_row_raises_on_missing_sheet(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    _build_multi_sheet_workbook(path)
    with pytest.raises(WorkbookParseError):
        find_header_row(path, "NoSuchSheet")


def test_parse_sheet_raises_on_all_blank_sheet(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empty"
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    with pytest.raises(WorkbookParseError):
        parse_sheet(path, "Empty")
