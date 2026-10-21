"""Offline contract test for the HCAI patient-origin/market-share adapter.

Fixture (pipelines/fixtures/hcai_patient_origin_sample.xlsx) is a real
3-sheet workbook subset: 2 rows where the facility is in Santa Clara
County, 2 rows where the patient's county of residence is Santa Clara
(facility elsewhere), and 1 row involving neither (for the filter test).
"""

from __future__ import annotations

from pathlib import Path

from scc_health_pipeline.sources.base import RawArtifact
from scc_health_pipeline.sources.hcai_patient_origin import HcaiPatientOriginAdapter

FIXTURES_DIR = Path(__file__).resolve().parents[1] / "fixtures"


def _fixture_artifact(tmp_path: Path) -> RawArtifact:
    raw_dir = tmp_path / "data" / "raw"
    raw_dir.mkdir(parents=True)
    fixture_copy = raw_dir / "patient_origin_sample.xlsx"
    fixture_copy.write_bytes((FIXTURES_DIR / "hcai_patient_origin_sample.xlsx").read_bytes())
    return RawArtifact(
        resource_id="fixture",
        source_id="hcai_patient_origin_market_share",
        local_path=fixture_copy,
        url="file://fixture",
        sha256="fixture",
        bytes=fixture_copy.stat().st_size,
        content_type="x",
        retrieved_at="2026-07-12T00:00:00+00:00",
        status="success",
    )


def test_normalize_keeps_rows_for_either_facility_or_patient_in_santa_clara(
    tmp_path: Path,
) -> None:
    adapter = HcaiPatientOriginAdapter()
    artifact = _fixture_artifact(tmp_path)
    outputs = adapter.normalize(artifact)
    assert len(outputs) == 1

    import polars as pl

    df = pl.read_parquet(outputs[0])
    assert df.height == 4  # the "neither" row is excluded
    assert "known_limitation" in df.columns
    assert "ambulatory surgery" in df["known_limitation"][0].lower()


def test_quality_checks_pass_on_fixture(tmp_path: Path) -> None:
    adapter = HcaiPatientOriginAdapter()
    artifact = _fixture_artifact(tmp_path)
    outputs = adapter.normalize(artifact)
    report = adapter.quality_checks(outputs)
    assert report.passed


def test_native_zip_geography_is_never_silently_converted_to_tract(tmp_path: Path) -> None:
    adapter = HcaiPatientOriginAdapter()
    artifact = _fixture_artifact(tmp_path)
    outputs = adapter.normalize(artifact)

    import polars as pl

    df = pl.read_parquet(outputs[0])
    assert "patient_zip" in df.columns
    assert "tract_geoid_2020" not in df.columns  # crosswalking is a separate, later step


def test_normalize_keeps_only_the_reporting_year_when_workbook_holds_several(
    tmp_path: Path,
) -> None:
    """HCAI now publishes 2024 and 2025 in one workbook; only 2024 rows may be
    labeled reporting_year=2024, so 2025 rows are never mixed in or double counted."""
    import openpyxl
    import polars as pl

    wb = openpyxl.Workbook()
    wb.active.title = "Patient Origin"
    wb.create_sheet("Market Share")
    data = wb.create_sheet("Data")
    data.append(
        [
            "oshpd_id",
            "pattype",
            "PZIP",
            "pcounty",
            "FACILITY_NAME",
            "COUNTY_NAME",
            "discharges",
            "year",
        ]
    )
    data.append(["1", "ED", "95112", "SANTA CLARA", "A HOSPITAL", "SANTA CLARA", 10, 2024])
    data.append(["1", "ED", "95112", "SANTA CLARA", "A HOSPITAL", "SANTA CLARA", 99, 2025])
    raw_dir = tmp_path / "data" / "raw"
    raw_dir.mkdir(parents=True)
    path = raw_dir / "multi_year.xlsx"
    wb.save(path)

    artifact = RawArtifact(
        resource_id="fixture",
        source_id="hcai_patient_origin_market_share",
        local_path=path,
        url="file://fixture",
        sha256="fixture",
        bytes=path.stat().st_size,
        content_type="x",
        retrieved_at="2026-10-20T00:00:00+00:00",
        status="success",
    )
    outputs = HcaiPatientOriginAdapter().normalize(artifact)
    df = pl.read_parquet(outputs[0])
    assert df.height == 1
    assert df["discharges"][0] == 10
    assert df["reporting_year"][0] == 2024
